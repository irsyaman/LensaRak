"""
quotation.py -- Generates a downloadable "Booking Summary / Quotation Draft"
(.xlsx) for a confirmed event.

CRITICAL RULE (same anti-hallucination pattern as sdg.py): the catalogue has
NO real monetary figures anywhere -- only a categorical Cost_Band (Low /
Medium / High) per offering. This module never invents a Ringgit amount.
Every price cell in the generated file is left BLANK for a human (the
organiser, using their real rate card) to fill in -- the Grand Total cell is
a live SUM() formula over those blanks, so the moment a human types real
numbers in, the total computes itself; until then it correctly shows 0.

Uses only openpyxl (already a project dependency) -- no new packages.
"""

from __future__ import annotations

import io
import sqlite3
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COST_BAND_NOTE = {
    "Low": "Low cost band (per catalogue) -- indicative only, not a price.",
    "Medium": "Medium cost band (per catalogue) -- indicative only, not a price.",
    "High": "High cost band (per catalogue) -- indicative only, not a price.",
}

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=16, bold=True)
WARN_FONT = Font(italic=True, color="B45309")


def _event_and_reservations(conn: sqlite3.Connection, event_id: int) -> tuple[dict | None, list[dict]]:
    ev = conn.execute("SELECT * FROM events WHERE event_id = ?", (event_id,)).fetchone()
    if ev is None:
        return None, []
    rows = conn.execute(
        "SELECT r.*, i.item_name, i.unit FROM reservations r JOIN items i ON i.item_code = r.item_code "
        "WHERE r.event_id = ? AND r.status = 'active'", (event_id,)
    ).fetchall()
    return dict(ev), [dict(r) for r in rows]


def build_quotation_xlsx(conn: sqlite3.Connection, event_id: int, offering: dict | None) -> bytes | None:
    """Returns the .xlsx file's bytes, or None if the event doesn't exist.
    `offering` is the Offerings_Master row dict for this event's offering_id
    (pass None if it can't be found -- the file still generates, just with
    blank offering-level fields instead of guessed ones)."""
    ev, reservations = _event_and_reservations(conn, event_id)
    if ev is None:
        return None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quotation Draft"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 16

    r = 1
    ws.cell(r, 1, "LensaRak -- Booking Summary / Quotation Draft").font = TITLE_FONT
    r += 1
    ws.cell(r, 1, "Generated " + datetime.now().strftime("%d %b %Y, %I:%M %p")).font = Font(italic=True, size=9)
    r += 2

    ws.cell(r, 1, "⚠ This is a booking summary, not a priced quotation.").font = WARN_FONT
    r += 1
    ws.cell(r, 1, "LensaRak's catalogue holds no Ringgit figures -- only a Low/Medium/High cost band. "
                  "Enter your organisation's real rates in the blank price cells below; "
                  "the Grand Total updates itself once you do.").font = Font(italic=True, size=9)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ws.row_dimensions[r].height = 28
    ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
    r += 2

    def kv(label, value):
        nonlocal r
        ws.cell(r, 1, label).font = Font(bold=True)
        ws.cell(r, 2, value if value not in (None, "") else "-")
        r += 1

    kv("Event ID", ev["event_id"])
    kv("Activity / Offering", ev.get("title") or (offering.get("Activity_Title") if offering else None))
    kv("Offering ID", ev.get("offering_id"))
    if offering:
        kv("Offering type", offering.get("Offering_Type"))
    kv("Date", ev["start_datetime"][:10])
    kv("Time", f"{ev['start_datetime'][11:]} - {ev['end_datetime'][11:]}")
    kv("Venue", ev.get("venue"))
    kv("Participants", ev.get("participants"))
    if offering:
        cost_band = offering.get("Cost_Band") or "Unknown"
        kv("Catalogue cost band", cost_band)
        ws.cell(r - 1, 3, COST_BAND_NOTE.get(cost_band, "Not recorded in catalogue.")).font = Font(italic=True, size=9)
    kv("Status", ev.get("status"))
    r += 1

    # ---- Reserved items table -------------------------------------------
    ws.cell(r, 1, "Reserved item").font = HEADER_FONT
    ws.cell(r, 2, "Quantity").font = HEADER_FONT
    ws.cell(r, 3, "Unit").font = HEADER_FONT
    ws.cell(r, 4, "Unit price (RM)").font = HEADER_FONT
    ws.cell(r, 5, "Line total (RM)").font = HEADER_FONT
    for c in range(1, 6):
        ws.cell(r, c).fill = HEADER_FILL
    header_row = r
    r += 1
    first_item_row = r
    if reservations:
        for res_row in reservations:
            ws.cell(r, 1, res_row["item_name"])
            ws.cell(r, 2, res_row["quantity_reserved"])
            ws.cell(r, 3, res_row.get("unit") or "")
            # price cells intentionally left BLANK -- never invented
            price_cell = f"D{r}"
            total_cell = f"{get_column_letter(5)}{r}"
            ws.cell(r, 5, f"=IF(D{r}=\"\",\"\",B{r}*D{r})")
            r += 1
    else:
        ws.cell(r, 1, "(No inventory items were reserved for this event -- e.g. a facilitator-led "
                       "session with no tracked equipment.)").font = Font(italic=True)
        r += 1
    last_item_row = r - 1

    r += 1
    ws.cell(r, 1, "Grand Total (RM)").font = Font(bold=True)
    if reservations:
        ws.cell(r, 5, f"=SUM(E{first_item_row}:E{last_item_row})").font = Font(bold=True)
    else:
        ws.cell(r, 5, 0).font = Font(bold=True)
    r += 2

    ws.cell(r, 1, "Prepared by (name / signature)"); r += 1
    ws.cell(r, 1, "Date"); r += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
