"""
loader.py -- Loads the organiser's Programme Catalogue (.xlsx) into plain
Python structures, plus small text-matching helpers shared by constraints.py
and matcher.py.

Expected sheets in the catalogue workbook:
  - Offerings_Master           one row per programme offering (ACT-001 ...)
  - Constraint_Rules           RULE-001 .. RULE-012
  - Theme_Objective_Mapping    MAP-001 ... (theme <-> stakeholder objective)
  - ACT-XXX_<name>             one sheet per offering, its bill of materials
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"


def find_default_catalogue_path() -> Path:
    """Looks for a few known names (including the organiser's original
    "DATASET_PROGRAMME CATALOGUE.xlsx", already sitting in data/ on this
    project), then falls back to a case-insensitive search for anything
    with "catalogue" in the name."""
    for name in ("programme_catalogue.xlsx", "DATASET_PROGRAMME CATALOGUE.xlsx"):
        candidate = DATA_DIR / name
        if candidate.exists():
            return candidate
    if DATA_DIR.exists():
        matches = sorted(p for p in DATA_DIR.glob("*.xlsx") if "catalogue" in p.name.lower())
        if matches:
            return matches[0]
    return DATA_DIR / "programme_catalogue.xlsx"


DEFAULT_CATALOGUE_PATH = find_default_catalogue_path()


# ---------------------------------------------------------------------------
# Small text-matching helpers (no NLP libraries -- crude token overlap, good
# enough for matching free-text theme/objective/BOM wording against the
# catalogue's own wording)
# ---------------------------------------------------------------------------

def _norm(s) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(s or "").lower()).strip()


def tokens(s) -> set[str]:
    return {t for t in _norm(s).split() if len(t) >= 2}


def token_variants(toks: set[str]) -> set[str]:
    """Adds a crude singular/plural variant of every token (e.g. 'wires' ->
    also matches 'wire') without needing a real stemmer."""
    out = set(toks)
    for t in toks:
        if t.endswith("s") and len(t) > 3:
            out.add(t[:-1])
        else:
            out.add(t + "s")
    return out


def token_overlap_ratio(a, b) -> float:
    ta, tb = tokens(a), tokens(b)
    if not ta or not tb:
        return 0.0
    tb_expanded = token_variants(tb)
    return len(ta & tb_expanded) / len(ta)


def parse_min_age(age_str) -> int | None:
    if not age_str:
        return None
    m = re.search(r"(\d+)", str(age_str))
    return int(m.group(1)) if m else None


def parse_number(s) -> float | None:
    if s is None or s == "":
        return None
    try:
        return float(s)
    except (TypeError, ValueError):
        m = re.search(r"(\d+(?:\.\d+)?)", str(s))
        return float(m.group(1)) if m else None


def _sheet_records(ws) -> list[dict]:
    """Turns a worksheet into a list of {header: value} dicts, using row 1
    as the header -- tolerant of blank leading rows/columns."""
    rows = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
    if not rows:
        return []
    header = rows[0]
    out = []
    for r in rows[1:]:
        d = {}
        for h, v in zip(header, r):
            if h is not None:
                d[str(h).strip()] = v
        if any(v is not None for v in d.values()):
            out.append(d)
    return out


# ---------------------------------------------------------------------------
# Catalogue loading
# ---------------------------------------------------------------------------

def load_catalogue(catalogue_path: Path = DEFAULT_CATALOGUE_PATH):
    """Returns (offerings, offerings_by_id, constraint_rules, theme_mapping,
    activity_items):
      offerings        -- list of dicts, one per row of Offerings_Master
      offerings_by_id  -- {Offering_ID: offering dict}
      constraint_rules -- list of dicts, one per row of Constraint_Rules
      theme_mapping     -- list of dicts, one per row of Theme_Objective_Mapping
      activity_items    -- {Offering_ID: [item text, ...]} from each ACT-XXX sheet
    """
    wb = openpyxl.load_workbook(catalogue_path, read_only=True, data_only=True)

    offerings = [d for d in _sheet_records(wb["Offerings_Master"]) if d.get("Offering_ID")]
    offerings_by_id = {d["Offering_ID"]: d for d in offerings}

    constraint_rules = _sheet_records(wb["Constraint_Rules"]) if "Constraint_Rules" in wb.sheetnames else []
    theme_mapping = _sheet_records(wb["Theme_Objective_Mapping"]) if "Theme_Objective_Mapping" in wb.sheetnames else []

    activity_items: dict[str, list[str]] = {}
    for name in wb.sheetnames:
        if not name.startswith("ACT-"):
            continue
        offering_id = name.split("_", 1)[0].strip()
        ws = wb[name]
        rows = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
        if not rows:
            continue
        header = rows[0]
        item_col = None
        for i, c in enumerate(header):
            if c and str(c).strip().lower().startswith("item"):
                item_col = i
                break
        if item_col is None:
            item_col = 2  # observed default position across most ACT-XXX sheets
        texts = []
        for r in rows[1:]:
            if item_col < len(r) and r[item_col]:
                texts.append(str(r[item_col]))
        activity_items[offering_id] = texts

    return offerings, offerings_by_id, constraint_rules, theme_mapping, activity_items


# ---------------------------------------------------------------------------
# Dropdown vocabulary for the web form -- pulled straight from the catalogue
# so pickers only ever suggest things Petrosains actually offers.
# ---------------------------------------------------------------------------

def _split_multi(val):
    return [p.strip() for p in str(val or "").split(";") if p.strip()]


def extract_choice_lists(offerings, theme_mapping):
    """Returns (themes, objectives, audiences, formats) -- all real values
    seen in the catalogue, sorted and de-duplicated."""
    themes = sorted({m["Theme"].strip() for m in theme_mapping if m.get("Theme")})
    objectives = sorted({m["Stakeholder_Objective"].strip() for m in theme_mapping if m.get("Stakeholder_Objective")})
    audiences = sorted({a for o in offerings for a in _split_multi(o.get("Audience_Types"))})
    formats = sorted({o["Offering_Type"].strip() for o in offerings if o.get("Offering_Type")})
    return themes, objectives, audiences, formats
