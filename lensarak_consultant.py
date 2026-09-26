#!/usr/bin/env python3
"""
LensaRak Consultant -- AI Programme Operations Consultant (Task 2 / Finals twist).

Recommends verified Petrosains programme offerings to a stakeholder (school,
public, corporate) based on their stated needs, using the organiser-supplied
DATASET_PROGRAMME CATALOGUE.xlsx as the single source of truth. Never invents
an offering, capacity, price, safety condition or availability -- everything
that can be checked against real data IS checked against real data; the LLM
(Gemini) is only used to turn an already-verified shortlist into readable
storyline/justification text, never to decide which offerings are feasible.

This module is completely separate from the Task 1 vision/checkout system
(lensarak_vision.py / lensarak_ops.py) -- it only *reads* the same inventory
database (lensarak_inventory.db), read-only, to check whether the materials
an offering needs are actually in stock right now.

USAGE
-----
    python lensarak_consultant.py                    interactive mode (asks questions)
    python lensarak_consultant.py --scenario a        canned demo: normal request
    python lensarak_consultant.py --scenario b        canned demo: missing information
    python lensarak_consultant.py --scenario c        canned demo: impossible / overflow request
    python lensarak_consultant.py --export            also saves the report to a .txt file
    python lensarak_consultant.py --gemini-key AIza...  enable AI storyline generation
                                  (or set the GEMINI_API_KEY environment variable)

Without a Gemini key the tool still works end-to-end (constraint filtering,
scoring, live inventory check, capacity/rotation planning) -- it just falls
back to a simpler templated storyline instead of an AI-generated one.

DATA FRESHNESS
--------------
The inventory database (lensarak_inventory.db) is normally kept at the
storeroom, which may not have the same network as wherever this tool is run
(e.g. an office upstairs). This tool does not try to sync itself -- it just
reads whatever copy of the .db file is sitting next to it, and always shows
you the file's last-modified time so you know how fresh that snapshot is.
Refreshing it is a manual step: copy the current lensarak_inventory.db from
the store machine over (USB drive, WhatsApp/Telegram to yourself, shared
Drive folder -- whichever is fastest) and replace the file in this folder.
"""

from __future__ import annotations

import argparse
import json
import math
import os
import re
import sqlite3
import sys
import textwrap
from datetime import datetime
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent
DEFAULT_DB_PATH = HERE / "lensarak_inventory.db"
DEFAULT_CATALOGUE_PATH = HERE / "data" / "DATASET_PROGRAMME CATALOGUE.xlsx"

GEMINI_API_KEY = os.environ.get("GEMINI_API_KEY", "")
GEMINI_MODEL = "gemini-2.5-flash"

COST_RANK = {"low": 1, "medium": 2, "high": 3}

SYSTEM_POLICY = (
    "You are the LensaRak AI Programme Operations Consultant for Petrosains. "
    "Recommend only from the supplied verified offerings list. Never invent an "
    "offering, duration, capacity, pricing, inventory quantity, safety condition "
    "or availability -- use only the values given to you. If required information "
    "is missing, say so or state the explicit assumption already given to you; "
    "never silently assume. Any idea you generate that is not one of the supplied "
    "verified offerings must be labelled 'Proposed Enhancement' and never presented "
    "as an existing Petrosains offering. Return structured JSON only, matching "
    "exactly the schema described in the prompt -- no prose outside the JSON."
)


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------

def _norm(s) -> str:
    """Lowercase, strip punctuation to spaces -- for loose text matching."""
    return re.sub(r"[^a-z0-9]+", " ", str(s or "").lower()).strip()


def _tokens(s) -> set[str]:
    return {t for t in _norm(s).split() if len(t) >= 2}


def _token_variants(tokens: set[str]) -> set[str]:
    """Adds a crude singular/plural variant of every token, so 'wires' still
    matches 'wire' etc. without a real stemmer."""
    out = set(tokens)
    for t in tokens:
        if t.endswith("s") and len(t) > 3:
            out.add(t[:-1])
        else:
            out.add(t + "s")
    return out


def _token_overlap_ratio(a, b) -> float:
    ta, tb = _tokens(a), _tokens(b)
    if not ta or not tb:
        return 0.0
    tb_expanded = _token_variants(tb)
    return len(ta & tb_expanded) / len(ta)


def _parse_min_age(age_str) -> int | None:
    if not age_str:
        return None
    m = re.search(r"(\d+)", str(age_str))
    return int(m.group(1)) if m else None


def _parse_number(s) -> float | None:
    if s is None or s == "":
        return None
    try:
        return float(s)
    except (TypeError, ValueError):
        m = re.search(r"(\d+(?:\.\d+)?)", str(s))
        return float(m.group(1)) if m else None


def _find_header_row(rows):
    """Returns (header_row_index, header_tuple) -- the first row containing
    a cell whose text isn't just None/blank, treated as the header."""
    for i, r in enumerate(rows):
        if any(c is not None for c in r):
            return i, r
    return 0, rows[0] if rows else ()


def _sheet_records(ws) -> list[dict]:
    rows = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
    if not rows:
        return []
    header = rows[0]
    out = []
    for r in rows[1:]:
        d = {}
        for h, v in zip(header, r):
            if h is not None:
                d[str(h).strip()] = v
        if any(v is not None for v in d.values()):
            out.append(d)
    return out


# ---------------------------------------------------------------------------
# Data loading -- catalogue (xlsx) + inventory (SQLite)
# ---------------------------------------------------------------------------

def load_catalogue(catalogue_path: Path):
    """Loads the organiser catalogue into plain Python structures.

    Returns (offerings, offerings_by_id, constraint_rules, theme_mapping,
    activity_items) where activity_items maps Offering_ID -> list[str] of
    every "Item" cell found in that offering's ACT-XXX sheet (its bill of
    materials), used later to cross-check against live inventory.
    """
    wb = openpyxl.load_workbook(catalogue_path, read_only=True, data_only=True)

    offerings = [d for d in _sheet_records(wb["Offerings_Master"]) if d.get("Offering_ID")]
    offerings_by_id = {d["Offering_ID"]: d for d in offerings}

    constraint_rules = _sheet_records(wb["Constraint_Rules"]) if "Constraint_Rules" in wb.sheetnames else []
    theme_mapping = _sheet_records(wb["Theme_Objective_Mapping"]) if "Theme_Objective_Mapping" in wb.sheetnames else []

    activity_items: dict[str, list[str]] = {}
    for name in wb.sheetnames:
        if not name.startswith("ACT-"):
            continue
        offering_id = name.split("_", 1)[0].strip()
        ws = wb[name]
        rows = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
        if not rows:
            continue
        header = rows[0]
        item_col = None
        for i, c in enumerate(header):
            if c and str(c).strip().lower().startswith("item"):
                item_col = i
                break
        if item_col is None:
            item_col = 2  # observed default position across most ACT-XXX sheets
        texts = []
        for r in rows[1:]:
            if item_col < len(r) and r[item_col]:
                texts.append(str(r[item_col]))
        activity_items[offering_id] = texts

    return offerings, offerings_by_id, constraint_rules, theme_mapping, activity_items


def load_inventory(db_path: Path) -> list[dict]:
    conn = sqlite3.connect(str(db_path))
    conn.row_factory = sqlite3.Row
    try:
        rows = conn.execute(
            "SELECT item_code, item_name, category, unit, total_quantity, "
            "available_quantity, status FROM items"
        ).fetchall()
    finally:
        conn.close()
    return [dict(r) for r in rows]


def db_freshness(db_path: Path) -> str:
    if not db_path.exists():
        return "TIADA FAIL -- lensarak_inventory.db tak dijumpai dalam folder ini"
    ts = datetime.fromtimestamp(db_path.stat().st_mtime)
    return ts.strftime("%d %b %Y, %I:%M %p")


# ---------------------------------------------------------------------------
# Inventory matching -- which real items does an offering actually need?
# ---------------------------------------------------------------------------

def match_offering_inventory(offering_id: str, activity_items: dict, inventory: list[dict]) -> list[dict]:
    texts = activity_items.get(offering_id, [])
    if not texts:
        return []
    blob_tokens = _token_variants(_tokens(" ".join(texts)))
    matched = []
    for item in inventory:
        name_tokens = _tokens(item["item_name"])
        if not name_tokens:
            continue
        overlap = name_tokens & blob_tokens
        if len(overlap) == len(name_tokens):
            matched.append(item)
        elif len(name_tokens) >= 3 and len(overlap) >= len(name_tokens) - 1:
            matched.append(item)
    return matched


def check_inventory_feasibility(offering_id: str, activity_items: dict, inventory: list[dict],
                                 participant_count: float | None) -> dict:
    matched = match_offering_inventory(offering_id, activity_items, inventory)
    if not matched:
        return {
            "status": "Unknown",
            "matched_items": [],
            "note": "Tiada padanan item inventory dikesan secara automatik dari senarai bahan aktiviti ini "
                    "-- sila semak stok secara manual.",
        }
    status = "Ready"
    details = []
    for it in matched:
        avail = it["available_quantity"] or 0
        if avail <= 0:
            status = "Unavailable"
        elif participant_count and avail < participant_count and status != "Unavailable":
            status = "Limited"
        details.append({
            "item_code": it["item_code"], "item_name": it["item_name"],
            "available_quantity": avail, "total_quantity": it["total_quantity"],
        })
    return {"status": status, "matched_items": details, "note": ""}


# ---------------------------------------------------------------------------
# Hard constraints (deterministic -- Constraint_Rules RULE-001..012)
# ---------------------------------------------------------------------------

def apply_hard_constraints(request: dict, offerings: list[dict]):
    feasible, excluded = [], []
    req_budget_rank = COST_RANK.get(str(request.get("budget") or "").strip().lower())

    for offering in offerings:
        exclude_reasons = []
        flags = []

        min_age = _parse_min_age(offering.get("Recommended_Age"))
        if request.get("age") is not None and min_age is not None and request["age"] < min_age:
            exclude_reasons.append(
                f"RULE-001 Umur: offering perlukan {offering.get('Recommended_Age')}, "
                f"permintaan menyatakan umur {request['age']}")
        elif request.get("age") is None and min_age is not None:
            flags.append(
                f"RULE-001 Umur BELUM disahkan: offering ini perlu umur minimum "
                f"{offering.get('Recommended_Age')} -- sahkan umur peserta sebelum ini dijadikan cadangan muktamad")

        if str(offering.get("Electricity_Required") or "").strip().lower() == "yes" \
                and str(request.get("electricity") or "").strip().lower() == "no":
            exclude_reasons.append("RULE-003 Elektrik: offering perlukan elektrik, venue dinyatakan tiada elektrik")

        if str(request.get("venue") or "").strip().lower() == "outdoor" \
                and str(offering.get("Indoor_Outdoor") or "").strip().lower() == "indoor":
            flags.append("RULE-006 Venue: offering ini disahkan untuk indoor sahaja -- sahkan keperluan "
                          "perlindungan/permukaan sebelum guna outdoor")

        setup = _parse_number(offering.get("Setup_Time_Min")) or 0
        std_dur = _parse_number(offering.get("Standard_Duration_Min")) or 0
        req_dur = _parse_number(request.get("duration_min"))
        if req_dur is not None and std_dur and req_dur < (setup + std_dur) * 0.6:
            flags.append(
                f"RULE-005 Tempoh: tempoh diminta ({int(req_dur)} min) mungkin tak cukup -- offering ini "
                f"biasanya perlukan ~{int(setup + std_dur)} min termasuk persediaan")

        off_cost_rank = COST_RANK.get(str(offering.get("Cost_Band") or "").strip().lower())
        if req_budget_rank and off_cost_rank:
            if off_cost_rank - req_budget_rank >= 2:
                exclude_reasons.append(
                    f"RULE-008 Bajet: cost band offering '{offering.get('Cost_Band')}' jauh melebihi "
                    f"bajet diminta")
            elif off_cost_rank > req_budget_rank:
                flags.append(f"RULE-008 Bajet: cost band '{offering.get('Cost_Band')}' lebih tinggi dari "
                              f"bajet diminta -- sila sahkan")

        max_p = _parse_number(offering.get("Max_Participants"))
        req_p = _parse_number(request.get("participant_count"))
        if req_p and max_p and req_p > max_p:
            rotations = math.ceil(req_p / max_p)
            flags.append(f"RULE-002 Kapasiti: max {int(max_p)} peserta/sesi -- perlukan {rotations} "
                         f"rotasi/sesi berasingan untuk {int(req_p)} peserta")

        if str(request.get("accessibility") or "").strip() and str(request.get("accessibility")).lower() not in ("no", "tiada", ""):
            flags.append(f"RULE-011 Aksesibiliti: {offering.get('Accessibility_Notes') or 'tiada nota khusus direkodkan'}")

        if exclude_reasons:
            excluded.append({**offering, "_exclude_reasons": exclude_reasons})
        else:
            feasible.append({**offering, "_flags": flags})

    return feasible, excluded


# ---------------------------------------------------------------------------
# Weighted scoring (0-100, transparent breakdown)
# ---------------------------------------------------------------------------

def score_offering(request: dict, offering: dict, theme_mapping: list[dict], resource_status: str) -> dict:
    theme_score = _token_overlap_ratio(request.get("theme"), offering.get("Suitable_Themes")) * 25
    for m in theme_mapping:
        if offering.get("Offering_ID") in (m.get("Primary_Offering_ID"), m.get("Secondary_Offering_ID")):
            bonus = _token_overlap_ratio(request.get("theme"), m.get("Theme")) * 25
            theme_score = max(theme_score, bonus)

    objective_score = _token_overlap_ratio(request.get("objective"), offering.get("Suitable_Objectives")) * 20
    for m in theme_mapping:
        if offering.get("Offering_ID") in (m.get("Primary_Offering_ID"), m.get("Secondary_Offering_ID")):
            bonus = _token_overlap_ratio(request.get("objective"), m.get("Stakeholder_Objective")) * 20
            objective_score = max(objective_score, bonus)

    audience = str(request.get("audience") or "").strip().lower()
    audience_types = str(offering.get("Audience_Types") or "").lower()
    if audience and audience in audience_types:
        audience_score = 15
    elif audience and _token_overlap_ratio(audience, audience_types) > 0:
        audience_score = 8
    else:
        audience_score = 5  # neutral -- not stated / not conclusive

    preferred_format = str(request.get("preferred_format") or "").strip().lower()
    offering_format = f"{offering.get('Offering_Type', '')} {offering.get('Delivery_Mode', '')}".lower()
    if preferred_format and preferred_format in offering_format:
        format_score = 10
    elif preferred_format:
        format_score = 3
    else:
        format_score = 5

    req_dur = _parse_number(request.get("duration_min"))
    std_dur = _parse_number(offering.get("Standard_Duration_Min"))
    if req_dur and std_dur:
        diff_ratio = abs(req_dur - std_dur) / std_dur
        duration_score = max(0.0, 10 * (1 - diff_ratio))
    else:
        duration_score = 5

    req_budget_rank = COST_RANK.get(str(request.get("budget") or "").strip().lower())
    off_cost_rank = COST_RANK.get(str(offering.get("Cost_Band") or "").strip().lower())
    if req_budget_rank and off_cost_rank:
        budget_score = 10 if req_budget_rank == off_cost_rank else (5 if abs(req_budget_rank - off_cost_rank) == 1 else 0)
    else:
        budget_score = 5

    resource_score = {"Ready": 10, "Limited": 5, "Unavailable": 0, "Unknown": 5}.get(resource_status, 5)

    # Don't let an age-gated, high-safety activity confidently rank as THE top
    # answer when the participants' age hasn't actually been confirmed yet --
    # it still appears (so nothing is hidden), just not presented with false
    # confidence ahead of safer, less age-critical options.
    age_penalty = 0.0
    min_age = _parse_min_age(offering.get("Recommended_Age"))
    if request.get("age") is None and min_age is not None and str(offering.get("Safety_Level", "")).strip().lower() == "high":
        age_penalty = -15.0

    breakdown = {
        "theme_alignment": round(theme_score, 1),
        "objective_alignment": round(objective_score, 1),
        "audience_alignment": round(audience_score, 1),
        "format_fit": round(format_score, 1),
        "duration_fit": round(duration_score, 1),
        "budget_fit": round(budget_score, 1),
        "resource_readiness": round(resource_score, 1),
    }
    if age_penalty:
        breakdown["age_unconfirmed_penalty"] = age_penalty
    total = round(sum(breakdown.values()), 1)
    return {"total": total, "breakdown": breakdown}


# ---------------------------------------------------------------------------
# Capacity overflow -- batch sessions vs multi-station rotation
# ---------------------------------------------------------------------------

def handle_capacity_overflow(request: dict, ranked: list[dict], activity_items: dict,
                              inventory: list[dict]) -> dict | None:
    req_p = _parse_number(request.get("participant_count"))
    if not req_p or not ranked:
        return None
    top = ranked[0]
    max_p = _parse_number(top["offering"].get("Max_Participants"))
    if not max_p or req_p <= max_p * 1.5:
        return None  # not a meaningful overflow -- ordinary rotation flag from constraints is enough

    # Option A -- batch / turn-based sessions using the top offering repeated
    rotations = math.ceil(req_p / max_p)
    option_a = {
        "label": "Sesi Berperingkat (Batch)",
        "description": f"Bahagikan {int(req_p)} peserta kepada {rotations} sesi berasingan "
                        f"(~{math.ceil(req_p / rotations)} peserta/sesi) menggunakan offering yang sama "
                        f"'{top['offering'].get('Activity_Title')}', dijalankan sesi demi sesi "
                        f"(pagi/petang atau merentas beberapa hari).",
        "rotations": rotations,
        "offering_id": top["offering"].get("Offering_ID"),
    }

    # Option B -- multi-station rotation using distinct verified offerings from
    # different STEM domains, running in parallel. Proposed Enhancement: this
    # specific rotation *arrangement* is our idea, even though every station
    # activity itself is a real verified offering.
    seen_domains = set()
    stations = []
    for cand in ranked:
        off = cand["offering"]
        domain_key = str(off.get("STEM_Domain") or "Lain-lain").split(";")[0].strip()
        if domain_key in seen_domains:
            continue
        seen_domains.add(domain_key)
        stations.append(off)
        if len(stations) == 4:
            break

    if len(stations) >= 2:
        n_stations = len(stations)
        per_station = math.ceil(req_p / n_stations)
        station_plan = []
        for off in stations:
            feas = check_inventory_feasibility(off.get("Offering_ID"), activity_items, inventory, per_station)
            station_plan.append({
                "offering_id": off.get("Offering_ID"),
                "title": off.get("Activity_Title"),
                "domain": off.get("STEM_Domain"),
                "participants_this_station": per_station,
                "resource_readiness": feas["status"],
            })
        option_b = {
            "label": "Proposed Enhancement",
            "name": "Model Rotasi Stesen (Multi-Station Rotation)",
            "description": f"Sedia {n_stations} stesen serentak, ~{per_station} peserta/stesen, setiap stesen "
                            f"guna offering verified yang berlainan domain supaya {int(req_p)} peserta boleh "
                            f"dilayan dalam SATU slot masa (bukan berulang-ulang).",
            "stations": station_plan,
        }
    else:
        option_b = None

    return {"option_a": option_a, "option_b": option_b}


# ---------------------------------------------------------------------------
# Request parsing (Gemini-assisted, with a plain-question fallback)
# ---------------------------------------------------------------------------

REQUEST_SCHEMA_FIELDS = [
    "theme", "objective", "audience", "age", "participant_count", "duration_min",
    "budget", "venue", "electricity", "internet", "water", "accessibility", "preferred_format",
]


def _get_gemini_client():
    if not GEMINI_API_KEY:
        return None
    try:
        from google import genai
        return genai.Client(api_key=GEMINI_API_KEY)
    except Exception as e:
        print(f"  (Gemini tak boleh digunakan -- {e}. Sambung tanpa AI storyline.)")
        return None


def parse_request_with_gemini(client, raw_text: str) -> dict | None:
    prompt = (
        "Extract structured fields from this Petrosains programme request. "
        "Return ONLY a JSON object with exactly these keys (null if not mentioned): "
        f"{REQUEST_SCHEMA_FIELDS}. "
        "age = a single representative age number if given. participant_count and duration_min = numbers only. "
        "budget = one of Low/Medium/High/null. electricity/internet/water = Yes/No/null. "
        "venue = Indoor/Outdoor/null.\n\n"
        f"Request: \"\"\"{raw_text}\"\"\""
    )
    try:
        response = client.models.generate_content(model=GEMINI_MODEL, contents=[prompt])
        text = (response.text or "").strip()
        text = re.sub(r"^```(json)?|```$", "", text.strip(), flags=re.MULTILINE).strip()
        data = json.loads(text)
        return {k: data.get(k) for k in REQUEST_SCHEMA_FIELDS}
    except Exception as e:
        print(f"  (Gemini gagal parse permintaan -- {e}. Guna soalan manual.)")
        return None


CRITICAL_FIELDS = ["age", "participant_count", "duration_min", "venue", "budget"]
FIELD_PROMPTS = {
    "theme": "Tema/tajuk program (contoh: tenaga boleh diperbaharui)",
    "objective": "Objektif stakeholder (contoh: kesedaran STEM)",
    "audience": "Jenis audiens (Primary/Secondary/Teachers/Public/Family)",
    "age": "Umur purata peserta (nombor)",
    "participant_count": "Bilangan peserta (nombor)",
    "duration_min": "Tempoh yang ada (minit)",
    "budget": "Bajet (Low/Medium/High)",
    "venue": "Venue (Indoor/Outdoor)",
    "electricity": "Elektrik tersedia? (Yes/No)",
    "internet": "Internet tersedia? (Yes/No)",
    "water": "Air tersedia? (Yes/No)",
    "accessibility": "Keperluan aksesibiliti khas? (tiada / nyatakan)",
    "preferred_format": "Format pilihan (Hands-on/Show/Demonstration/Challenge) -- boleh kosong",
}


def parse_request_core(raw_text: str, gemini_client) -> dict:
    """Parses free text into structured fields (Gemini-assisted, with a
    heuristic regex fallback so it still works with no API key). Does NOT
    resolve missing critical fields -- that's the caller's job (interactive
    prompt for the CLI, form widgets for the web dashboard, or an explicit
    stated assumption), so this function is shared by both front-ends."""
    request = {"raw_text": raw_text}
    parsed = parse_request_with_gemini(gemini_client, raw_text) if gemini_client else None
    if parsed:
        request.update(parsed)
    else:
        m = re.search(r"(\d+)\s*(pelajar|students?|peserta|participants?|orang)", raw_text, re.I)
        if m:
            request["participant_count"] = float(m.group(1))
        m = re.search(r"(\d+)\s*(minit|minutes?|min\b)", raw_text, re.I)
        if m:
            request["duration_min"] = float(m.group(1))
        m = re.search(r"(\d+)\s*(jam|hours?)", raw_text, re.I)
        if m and "duration_min" not in request:
            request["duration_min"] = float(m.group(1)) * 60
        request.setdefault("theme", raw_text)
        request.setdefault("objective", raw_text)

    for k in REQUEST_SCHEMA_FIELDS:
        request.setdefault(k, None)
    for k in ("age", "participant_count", "duration_min"):
        request[k] = _parse_number(request.get(k))
    return request


def missing_critical_fields(request: dict) -> list[str]:
    return [f for f in CRITICAL_FIELDS if request.get(f) in (None, "")]


def parse_request(raw_text: str, assume: bool, gemini_client) -> tuple[dict, list[str], list[str]]:
    """CLI front-end: resolves missing critical fields via an interactive
    prompt (or an explicit stated assumption with --assume / no tty).
    Returns (request, missing_information, assumptions)."""
    request = parse_request_core(raw_text, gemini_client)
    missing_information, assumptions = [], []
    for field in missing_critical_fields(request):
        if assume:
            assumptions.append(f"{FIELD_PROMPTS[field]}: tidak dinyatakan -- andaian dibuat sebagai 'tidak diketahui'")
            continue
        if sys.stdin.isatty():
            answer = input(f"  ? {FIELD_PROMPTS[field]}: ").strip()
            if answer:
                request[field] = _parse_number(answer) if field in ("age", "participant_count", "duration_min") else answer
                continue
        missing_information.append(FIELD_PROMPTS[field])

    return request, missing_information, assumptions


# ---------------------------------------------------------------------------
# Programme composer (Gemini narrative, verified-context only) + fallback
# ---------------------------------------------------------------------------

def compose_programme_llm(client, request, ranked, capacity_plan, missing_information, assumptions) -> dict | None:
    context = {
        "request": {k: v for k, v in request.items() if k != "raw_text"},
        "raw_request_text": request.get("raw_text"),
        "missing_information": missing_information,
        "assumptions": assumptions,
        "verified_candidates": [
            {
                "offering_id": c["offering"].get("Offering_ID"),
                "title": c["offering"].get("Activity_Title"),
                "short_description": c["offering"].get("Short_Description"),
                "match_score": c["score"]["total"],
                "score_breakdown": c["score"]["breakdown"],
                "flags": c["offering"].get("_flags", []),
                "resource_readiness": c["resource"]["status"],
            }
            for c in ranked[:5]
        ],
        "capacity_plan": capacity_plan,
    }
    prompt = (
        f"{SYSTEM_POLICY}\n\n"
        "Context (verified data -- do not add offerings not listed here):\n"
        f"{json.dumps(context, ensure_ascii=False, default=str)}\n\n"
        "Return JSON with exactly these top-level keys: "
        "programme (object: title, theme, objective, storyline), "
        "recommended_offerings (array of {offering_id, reason}), "
        "participant_journey (array of strings, ordered steps), "
        "proposed_enhancements (array of {name, description} -- label implied, creative ideas only), "
        "risks (array of strings), "
        "alternative_options (array of strings)."
    )
    try:
        response = client.models.generate_content(model=GEMINI_MODEL, contents=[prompt])
        text = (response.text or "").strip()
        text = re.sub(r"^```(json)?|```$", "", text.strip(), flags=re.MULTILINE).strip()
        return json.loads(text)
    except Exception as e:
        print(f"  (Gemini gagal hasilkan storyline -- {e}. Guna templat asas.)")
        return None


def compose_programme_fallback(request, ranked, capacity_plan) -> dict:
    if not ranked:
        return {
            "programme": {"title": "Tiada offering sesuai dijumpai", "theme": request.get("theme"),
                           "objective": request.get("objective"), "storyline": ""},
            "recommended_offerings": [], "participant_journey": [],
            "proposed_enhancements": [], "risks": ["Tiada offering verified lepas tapisan kekangan."],
            "alternative_options": [],
        }
    top = ranked[0]["offering"]
    theme_raw = request.get("theme") or "STEM"
    theme = theme_raw if len(theme_raw) <= 40 else theme_raw[:37].rsplit(" ", 1)[0] + "..."
    storyline = (
        f"Program bertemakan {theme} bermula dengan pengenalan konsep, diikuti aktiviti utama "
        f"'{top.get('Activity_Title')}' secara hands-on, dan ditutup dengan sesi refleksi/perkongsian hasil."
    )
    journey = [
        "Pengenalan & ice-breaker (5-10 min)",
        f"Aktiviti utama: {top.get('Activity_Title')} -- {top.get('Short_Description')}",
        "Sesi refleksi & perkongsian hasil (10 min)",
    ]
    if capacity_plan:
        journey.append("Nota operasi: lihat pelan rotasi/stesen di bawah untuk kumpulan besar.")
    return {
        "programme": {"title": f"Program {theme}", "theme": theme, "objective": request.get("objective"),
                       "storyline": storyline},
        "recommended_offerings": [{"offering_id": c["offering"].get("Offering_ID"),
                                    "reason": f"Skor padanan {c['score']['total']}/100 -- lihat pecahan skor."}
                                   for c in ranked[:3]],
        "participant_journey": journey,
        "proposed_enhancements": [],
        "risks": ["Storyline ini templat asas (tiada Gemini API key) -- set GEMINI_API_KEY untuk cerita "
                  "yang lebih kaya & cadangan enhancement."],
        "alternative_options": [f"{c['offering'].get('Activity_Title')} ({c['offering'].get('Offering_ID')})"
                                 for c in ranked[1:4]],
    }


def verify_output(programme: dict, offerings_by_id: dict, resource_by_id: dict, flags_by_id: dict) -> dict:
    """Anti-hallucination pass: overwrite any offering-linked fact with the
    real catalogue/inventory truth, and drop any offering_id that isn't real."""
    verified_recs = []
    risks = list(programme.get("risks") or [])
    for rec in programme.get("recommended_offerings") or []:
        oid = rec.get("offering_id")
        truth = offerings_by_id.get(oid)
        if not truth:
            risks.append(f"AI cuba cadangkan offering_id '{oid}' yang tidak wujud dalam katalog -- dibuang.")
            continue
        verified_recs.append({
            "offering_id": oid,
            "title": truth.get("Activity_Title"),
            "reason": rec.get("reason", ""),
            "duration_min": truth.get("Standard_Duration_Min"),
            "max_participants": truth.get("Max_Participants"),
            "cost_band": truth.get("Cost_Band"),
            "safety_level": truth.get("Safety_Level"),
            "recommended_age": truth.get("Recommended_Age"),
            "resource_readiness": resource_by_id.get(oid, {}).get("status", "Unknown"),
            "flags": flags_by_id.get(oid, []),
        })
    enhancements = []
    for e in programme.get("proposed_enhancements") or []:
        if isinstance(e, str):
            e = {"name": e, "description": ""}
        e["label"] = "Proposed Enhancement"
        enhancements.append(e)

    return {
        "programme": programme.get("programme", {}),
        "recommended_offerings": verified_recs,
        "participant_journey": programme.get("participant_journey") or [],
        "proposed_enhancements": enhancements,
        "risks": risks,
        "alternative_options": programme.get("alternative_options") or [],
    }


# ---------------------------------------------------------------------------
# Report printing / export
# ---------------------------------------------------------------------------

def print_report(request, missing_information, assumptions, excluded, verified, capacity_plan,
                  db_fresh_str, catalogue_fresh_str):
    W = 78
    def hr(title=""):
        print("=" * W if not title else f"\n{'=' * 3} {title} {'=' * max(0, W - 5 - len(title))}")

    hr("LENSARAK AI PROGRAMME CONSULTANT")
    print(f"Data inventory terakhir dikemaskini: {db_fresh_str}")
    print(f"Katalog program terakhir dikemaskini: {catalogue_fresh_str}")

    hr("1. PERMINTAAN STAKEHOLDER")
    print(textwrap.fill(request.get("raw_text") or "(tiada teks asal)", W))
    for k in REQUEST_SCHEMA_FIELDS:
        if request.get(k) not in (None, ""):
            print(f"  - {k}: {request[k]}")

    if missing_information:
        hr("2. MAKLUMAT HILANG / SOALAN SUSULAN")
        for m in missing_information:
            print(f"  ? {m}")
    if assumptions:
        hr("3. ANDAIAN")
        for a in assumptions:
            print(f"  * {a}")

    if excluded:
        hr("OFFERING DITOLAK (kekangan keras)")
        for e in excluded[:8]:
            print(f"  - {e.get('Activity_Title')} ({e.get('Offering_ID')})")
            for r in e["_exclude_reasons"]:
                print(f"      x {r}")

    prog = verified["programme"]
    hr("4. CADANGAN PROGRAM")
    print(f"Tajuk: {prog.get('title')}")
    print(textwrap.fill(prog.get("storyline") or "", W))

    hr("5. OFFERING DISYORKAN")
    for r in verified["recommended_offerings"]:
        print(f"  [{r['resource_readiness'].upper()}] {r['title']} ({r['offering_id']}) -- {r['reason']}")
        print(f"      Tempoh piawai: {r['duration_min']} min | Max peserta: {r['max_participants']} | "
              f"Cost band: {r['cost_band']} | Safety: {r['safety_level']} | Umur: {r['recommended_age']}")
        for f in r.get("flags") or []:
            print(f"      ! {f}")

    if capacity_plan:
        hr("6. PELAN KAPASITI (kumpulan besar)")
        a = capacity_plan["option_a"]
        print(f"  Pilihan A -- {a['label']}: {a['description']}")
        if capacity_plan.get("option_b"):
            b = capacity_plan["option_b"]
            print(f"  Pilihan B -- [{b['label']}] {b['name']}: {b['description']}")
            for s in b["stations"]:
                print(f"      Stesen: {s['title']} ({s['domain']}) -- {s['participants_this_station']} peserta "
                      f"-- stok: {s['resource_readiness']}")

    hr("7. PARTICIPANT JOURNEY")
    for i, step in enumerate(verified["participant_journey"], 1):
        print(f"  {i}. {step}")

    if verified["proposed_enhancements"]:
        hr("8. PROPOSED ENHANCEMENT (idea AI -- BUKAN offering rasmi)")
        for e in verified["proposed_enhancements"]:
            print(f"  * [{e['label']}] {e.get('name')}: {e.get('description', '')}")

    if verified["alternative_options"]:
        hr("9. ALTERNATIF")
        for a in verified["alternative_options"]:
            print(f"  - {a}")

    if verified["risks"]:
        hr("10. RISIKO / NOTA")
        for r in verified["risks"]:
            print(f"  ! {r}")
    hr()


def export_report_text(path: Path, *args):
    import io
    import contextlib
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        print_report(*args)
    path.write_text(buf.getvalue(), encoding="utf-8")
    print(f"\n(Report disimpan ke {path})")


# ---------------------------------------------------------------------------
# Demo scenarios (from the team's 7-hour implementation framework, section 11)
# ---------------------------------------------------------------------------

SCENARIOS = {
    "a": dict(raw_text="Kami perlukan program tenaga boleh diperbaharui untuk 30 pelajar sekolah rendah, 90 minit.",
              theme="Renewable Energy", objective="Sustainability awareness", audience="Primary",
              age=10, participant_count=30, duration_min=90, budget="Medium", venue="Indoor",
              electricity="Yes", internet=None, water=None, accessibility=None, preferred_format=None),
    "b": dict(raw_text="Kami mahukan program robotik.",
              theme="Robotics", objective=None, audience=None, age=None, participant_count=None,
              duration_min=None, budget=None, venue=None, electricity=None, internet=None, water=None,
              accessibility=None, preferred_format=None),
    "c": dict(raw_text="Kami perlukan aktiviti bina mini-dron individu untuk 500 pelajar prasekolah dalam 30 minit "
                       "dengan bajet sangat rendah.",
              theme="Aviation", objective="Engineering exposure", audience="Early Years", age=4,
              participant_count=500, duration_min=30, budget="Low", venue="Indoor", electricity="Yes",
              internet=None, water=None, accessibility=None, preferred_format=None),
}


# ---------------------------------------------------------------------------
# Main pipeline
# ---------------------------------------------------------------------------

def run(request: dict, missing_information: list[str], assumptions: list[str],
        offerings, offerings_by_id, theme_mapping, activity_items, inventory, gemini_client):
    feasible, excluded = apply_hard_constraints(request, offerings)

    scored = []
    for offering in feasible:
        resource = check_inventory_feasibility(offering.get("Offering_ID"), activity_items, inventory,
                                                 request.get("participant_count"))
        score = score_offering(request, offering, theme_mapping, resource["status"])
        scored.append({"offering": offering, "score": score, "resource": resource})
    scored.sort(key=lambda c: c["score"]["total"], reverse=True)

    resource_by_id = {c["offering"].get("Offering_ID"): c["resource"] for c in scored}
    flags_by_id = {c["offering"].get("Offering_ID"): c["offering"].get("_flags", []) for c in scored}
    capacity_plan = handle_capacity_overflow(request, scored, activity_items, inventory)

    llm_result = compose_programme_llm(gemini_client, request, scored, capacity_plan,
                                        missing_information, assumptions) if gemini_client else None
    programme = llm_result or compose_programme_fallback(request, scored, capacity_plan)
    verified = verify_output(programme, offerings_by_id, resource_by_id, flags_by_id)

    return excluded, scored, capacity_plan, verified


def main():
    ap = argparse.ArgumentParser(description="LensaRak AI Programme Operations Consultant")
    ap.add_argument("--db", type=Path, default=DEFAULT_DB_PATH, help="Path to lensarak_inventory.db")
    ap.add_argument("--catalogue", type=Path, default=DEFAULT_CATALOGUE_PATH,
                    help="Path to DATASET_PROGRAMME CATALOGUE.xlsx")
    ap.add_argument("--gemini-key", default=None, help="Gemini API key (or set GEMINI_API_KEY env var)")
    ap.add_argument("--scenario", choices=["a", "b", "c"], help="Run a canned demo scenario instead of asking")
    ap.add_argument("--assume", action="store_true",
                    help="Never prompt interactively -- state assumptions for any missing field instead")
    ap.add_argument("--export", action="store_true", help="Also save the report as a .txt file")
    args = ap.parse_args()

    global GEMINI_API_KEY
    if args.gemini_key:
        GEMINI_API_KEY = args.gemini_key
    gemini_client = _get_gemini_client()

    print("Memuatkan katalog program & inventori...")
    offerings, offerings_by_id, constraint_rules, theme_mapping, activity_items = load_catalogue(args.catalogue)
    inventory = load_inventory(args.db) if args.db.exists() else []
    if not inventory:
        print(f"  AMARAN: {args.db} tak dijumpai/kosong -- semakan stok akan menunjukkan 'Unknown' untuk semua.")

    if args.scenario:
        preset = SCENARIOS[args.scenario]
        request = {"raw_text": preset["raw_text"]}
        request.update({k: preset[k] for k in REQUEST_SCHEMA_FIELDS})
        missing_information, assumptions = [], []
        for field in CRITICAL_FIELDS:
            if request.get(field) in (None, ""):
                assumptions.append(f"{FIELD_PROMPTS[field]}: tidak dinyatakan dalam skrip demo -- andaian 'tidak diketahui'")
        print(f"\n[Mod demo -- skrip {args.scenario.upper()}]\n\"{preset['raw_text']}\"\n")
    else:
        print("\nTerangkan keperluan program stakeholder (satu ayat):")
        raw_text = input("> ").strip()
        request, missing_information, assumptions = parse_request(raw_text, args.assume, gemini_client)

    excluded, scored, capacity_plan, verified = run(
        request, missing_information, assumptions, offerings, offerings_by_id,
        theme_mapping, activity_items, inventory, gemini_client)

    db_fresh = db_freshness(args.db)
    cat_fresh = db_freshness(args.catalogue)

    print_report(request, missing_information, assumptions, excluded, verified, capacity_plan, db_fresh, cat_fresh)

    if args.export:
        out_path = HERE / f"lensarak_consultant_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt"
        export_report_text(out_path, request, missing_information, assumptions, excluded, verified,
                            capacity_plan, db_fresh, cat_fresh)


if __name__ == "__main__":
    main()
